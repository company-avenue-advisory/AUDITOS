"""
Regression tests for services.gstr1_generator.build_gstr1_excel_rows and
services.gstr1_excel_export.write_gstr1_excel - the GSTN offline-tool
Excel format (b2b, cdnr, hsn (b2b), hsn (b2c), b2cs, docs sheets), which
previously didn't exist at all (only the direct-API JSON format did).
"""
import sys
import os
import tempfile
import unittest
from dataclasses import dataclass
from typing import Optional

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "../../..")))

from backend.services.gstr1_generator import build_gstr1_excel_rows
from backend.services.gstr1_excel_export import write_gstr1_excel, _SHEET_COLUMNS, HEADER_ROW, DATA_START_ROW


@dataclass
class FakeLineItem:
    invoice_no: str
    party_gstin: Optional[str]
    gstr1_category: Optional[str]
    taxable_value: float
    cgst_amount: float = 0.0
    sgst_amount: float = 0.0
    igst_amount: float = 0.0
    total_invoice_value: float = 0.0
    voucher_date: str = "30-06-2026"
    voucher_type: str = "Sales"
    place_of_supply: Optional[str] = None
    hsn: Optional[str] = "9971"
    particulars: str = "SaaS/UPI Platform Charges"
    party_ledger_name: str = "Acme Bank"
    qty: float = 1.0


FIRM_GSTIN = "27AAAAA1111A1Z1"


class TestBuildGstr1ExcelRows(unittest.TestCase):

    def test_b2b_invoice_lands_on_b2b_sheet(self):
        items = [FakeLineItem("MH1", "27CCCCC3333C3Z3", "B2B", 1000.0,
                               cgst_amount=90.0, sgst_amount=90.0, total_invoice_value=1180.0)]
        rows = build_gstr1_excel_rows(items, FIRM_GSTIN)
        self.assertEqual(len(rows["b2b"]), 1)
        self.assertEqual(rows["b2b"][0]["gstin"], "27CCCCC3333C3Z3")
        self.assertEqual(rows["b2b"][0]["taxable_value"], 1000.0)
        self.assertEqual(rows["cdnr"], [])

    def test_credit_note_lands_on_cdnr_sheet_not_b2b(self):
        items = [FakeLineItem("CR1", "27CCCCC3333C3Z3", None, 500.0,
                               igst_amount=90.0, total_invoice_value=590.0,
                               voucher_type="Credit Note")]
        rows = build_gstr1_excel_rows(items, FIRM_GSTIN)
        self.assertEqual(rows["b2b"], [])
        self.assertEqual(len(rows["cdnr"]), 1)
        self.assertEqual(rows["cdnr"][0]["note_type"], "Credit Note")
        self.assertEqual(rows["cdnr"][0]["note_number"], "CR1")

    def test_debit_note_labeled_correctly_on_cdnr_sheet(self):
        items = [FakeLineItem("DR1", "27CCCCC3333C3Z3", None, 500.0,
                               igst_amount=90.0, total_invoice_value=590.0,
                               voucher_type="Debit Note")]
        rows = build_gstr1_excel_rows(items, FIRM_GSTIN)
        self.assertEqual(rows["cdnr"][0]["note_type"], "Debit Note")

    def test_b2c_unregistered_invoice_lands_on_b2cs_sheet(self):
        items = [FakeLineItem("MH2", None, "B2CS", 2000.0,
                               cgst_amount=180.0, sgst_amount=180.0, total_invoice_value=2360.0,
                               place_of_supply="27")]
        rows = build_gstr1_excel_rows(items, FIRM_GSTIN)
        self.assertEqual(len(rows["b2cs"]), 1)
        self.assertEqual(rows["b2cs"][0]["taxable_value"], 2000.0)

    def test_hsn_split_b2b_vs_b2c_by_category_not_doc_type(self):
        items = [
            FakeLineItem("MH1", "27CCCCC3333C3Z3", "B2B", 1000.0, hsn="9971",
                         cgst_amount=90.0, sgst_amount=90.0, total_invoice_value=1180.0),
            FakeLineItem("MH2", None, "B2CS", 500.0, hsn="997319",
                         cgst_amount=45.0, sgst_amount=45.0, total_invoice_value=590.0),
        ]
        rows = build_gstr1_excel_rows(items, FIRM_GSTIN)
        self.assertEqual(len(rows["hsn_b2b"]), 1)
        self.assertEqual(rows["hsn_b2b"][0]["hsn"], "9971")
        self.assertEqual(len(rows["hsn_b2c"]), 1)
        self.assertEqual(rows["hsn_b2c"][0]["hsn"], "997319")

    def test_hsn_credit_note_subtracts_not_omits(self):
        # Same fix as _build_hsn (JSON path) - a credit note must reduce its
        # HSN bucket's total, not vanish from the HSN sheet entirely.
        items = [
            FakeLineItem("MH1", "27CCCCC3333C3Z3", "B2B", 1000.0, hsn="9971",
                         cgst_amount=90.0, sgst_amount=90.0, total_invoice_value=1180.0),
            FakeLineItem("CR1", "27CCCCC3333C3Z3", None, 300.0, hsn="9971",
                         cgst_amount=27.0, sgst_amount=27.0, total_invoice_value=354.0,
                         voucher_type="Credit Note"),
        ]
        rows = build_gstr1_excel_rows(items, FIRM_GSTIN)
        self.assertEqual(len(rows["hsn_b2b"]), 1)
        self.assertEqual(rows["hsn_b2b"][0]["taxable_value"], 700.0)

    def test_docs_sheet_counts_by_nature(self):
        items = [
            FakeLineItem("MH1", "27X", "B2B", 1000.0),
            FakeLineItem("MH2", "27X", "B2B", 1000.0),
            FakeLineItem("CR1", "27X", None, 500.0, voucher_type="Credit Note"),
        ]
        rows = build_gstr1_excel_rows(items, FIRM_GSTIN)
        by_nature = {r["nature_of_document"]: r for r in rows["docs"]}
        self.assertEqual(by_nature["Invoices for outward supply"]["total_number"], 2)
        self.assertEqual(by_nature["Credit Note"]["total_number"], 1)
        self.assertNotIn("Debit Note", by_nature)


class TestWriteGstr1Excel(unittest.TestCase):

    def test_writes_a_workbook_with_all_six_sheets(self):
        items = [
            FakeLineItem("MH1", "27CCCCC3333C3Z3", "B2B", 1000.0, hsn="9971",
                         cgst_amount=90.0, sgst_amount=90.0, total_invoice_value=1180.0),
            FakeLineItem("MH2", None, "B2CS", 500.0, hsn="997319",
                         cgst_amount=45.0, sgst_amount=45.0, total_invoice_value=590.0),
            FakeLineItem("CR1", "27CCCCC3333C3Z3", None, 300.0, hsn="9971",
                         cgst_amount=27.0, sgst_amount=27.0, total_invoice_value=354.0,
                         voucher_type="Credit Note"),
        ]
        excel_rows = build_gstr1_excel_rows(items, FIRM_GSTIN)

        tmp_dir = tempfile.mkdtemp(prefix="gstr1_excel_test_")
        path = os.path.join(tmp_dir, "gstr1.xlsx")
        write_gstr1_excel(excel_rows, path)

        import openpyxl
        wb = openpyxl.load_workbook(path)
        # The real template has ~30 sheets total (b2cl, exp, at, master, ...);
        # we only populate these 6, the rest pass through untouched.
        self.assertTrue(set(_SHEET_COLUMNS.keys()).issubset(set(wb.sheetnames)))

        b2b_ws = wb["b2b"]
        header = [c.value for c in next(b2b_ws.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW))]
        self.assertEqual(header[0], "GSTIN/UIN of Recipient")
        self.assertEqual(b2b_ws.cell(row=DATA_START_ROW, column=1).value, "27CCCCC3333C3Z3")
        # Invoice date must land as a real Excel date, not a "DD-MM-YYYY" string
        # (a text value in this template's date-formatted column is what was
        # confirmed causing ExpressGST to mis-read/reject the workbook).
        import datetime
        date_col = header.index("Invoice date") + 1
        self.assertIsInstance(b2b_ws.cell(row=DATA_START_ROW, column=date_col).value, datetime.date)

    def test_empty_rows_still_write_headers_only(self):
        tmp_dir = tempfile.mkdtemp(prefix="gstr1_excel_empty_test_")
        path = os.path.join(tmp_dir, "empty.xlsx")
        write_gstr1_excel({}, path)

        import openpyxl
        wb = openpyxl.load_workbook(path)
        self.assertTrue(set(_SHEET_COLUMNS.keys()).issubset(set(wb.sheetnames)))
        self.assertIsNone(wb["b2b"].cell(row=DATA_START_ROW, column=1).value)

    def test_hsn_sheet_columns_are_not_swapped(self):
        # Regression: hsn (b2b)/(b2c) previously had Rate and Taxable Value
        # swapped relative to the real template because columns were
        # hardcoded by position instead of looked up by the template's own
        # header label.
        items = [FakeLineItem("MH1", "27CCCCC3333C3Z3", "B2B", 700.0, hsn="9971",
                               cgst_amount=63.0, sgst_amount=63.0, total_invoice_value=826.0)]
        excel_rows = build_gstr1_excel_rows(items, FIRM_GSTIN)
        tmp_dir = tempfile.mkdtemp(prefix="gstr1_excel_hsn_test_")
        path = os.path.join(tmp_dir, "gstr1.xlsx")
        write_gstr1_excel(excel_rows, path)

        import openpyxl
        wb = openpyxl.load_workbook(path)
        ws = wb["hsn (b2b)"]
        header = [c.value for c in next(ws.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW))]
        taxable_col = header.index("Taxable Value") + 1
        rate_col = header.index("Rate") + 1
        self.assertEqual(ws.cell(row=DATA_START_ROW, column=taxable_col).value, 700.0)
        self.assertEqual(ws.cell(row=DATA_START_ROW, column=rate_col).value, 18)


if __name__ == "__main__":
    unittest.main()
