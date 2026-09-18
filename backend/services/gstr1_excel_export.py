"""
gstr1_excel_export.py -- Writes a GSTR-1 workbook by filling data into
services/templates/gstr1_offline_tool_template.xlsx (sheets: b2b, cdnr,
hsn (b2b), hsn (b2c), b2cs, docs, plus the untouched official sheets we
don't populate), for accountants who file through the offline utility or
a third-party tool like ExpressGST instead of a direct GSP/API upload.

Previously this built a workbook from scratch with hand-typed column
lists, which drifted from the real template in three separate ways
(hsn (b2b)/(b2c) had Rate and Taxable Value swapped, docs had an extra
"Net Issued" column that doesn't exist in the real template, and Invoice
date/Note Date were written as plain text instead of the real numeric
Excel date the template's own cells are pre-formatted for) -- confirmed
as the cause of ExpressGST showing wrong/corrupted data on upload. Basing
the output on the actual template file makes that whole class of drift
structurally impossible: columns are located by matching the LABEL
against the template's own header row, not by a hardcoded position.

Input is services.gstr1_generator.build_gstr1_excel_rows(items, firm_gstin)'s
flat-rows dict - this file only handles the openpyxl mechanics.
"""
import logging
import os
from datetime import datetime
from typing import List

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    logging.warning("openpyxl not installed. Install it to enable GSTR-1 Excel export.")

logger = logging.getLogger(__name__)

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "templates", "gstr1_offline_tool_template.xlsx")

# The template's fixed layout: row 1 = section title, row 2 = live SUM
# formulas over the data range, row 3 = column headers, row 4+ = data.
HEADER_ROW = 3
DATA_START_ROW = 4

# (sheet name, [(row key, column label), ...]) - label must match the
# template's own header row exactly; column position is looked up from
# the template at write time, not assumed from this list's order.
_SHEET_COLUMNS = {
    "b2b": [
        ("gstin", "GSTIN/UIN of Recipient"),
        ("receiver_name", "Receiver Name"),
        ("invoice_number", "Invoice Number"),
        ("invoice_date", "Invoice date"),
        ("invoice_value", "Invoice Value"),
        ("place_of_supply", "Place Of Supply"),
        ("reverse_charge", "Reverse Charge"),
        ("applicable_tax_rate", "Applicable % of Tax Rate"),
        ("invoice_type", "Invoice Type"),
        ("ecommerce_gstin", "E-Commerce GSTIN"),
        ("rate", "Rate"),
        ("taxable_value", "Taxable Value"),
        ("cess", "Cess Amount"),
    ],
    "cdnr": [
        ("gstin", "GSTIN/UIN of Recipient"),
        ("receiver_name", "Receiver Name"),
        ("note_number", "Note Number"),
        ("note_date", "Note Date"),
        ("note_type", "Note Type"),
        ("place_of_supply", "Place Of Supply"),
        ("reverse_charge", "Reverse Charge"),
        ("note_supply_type", "Note Supply Type"),
        ("note_value", "Note Value"),
        ("applicable_tax_rate", "Applicable % of Tax Rate"),
        ("rate", "Rate"),
        ("taxable_value", "Taxable Value"),
        ("cess", "Cess Amount"),
    ],
    "hsn (b2b)": [
        ("hsn", "HSN"),
        ("description", "Description"),
        ("uqc", "UQC"),
        ("total_quantity", "Total Quantity"),
        ("total_value", "Total Value"),
        ("taxable_value", "Taxable Value"),
        ("rate", "Rate"),
        ("integrated_tax", "Integrated Tax Amount"),
        ("central_tax", "Central Tax Amount"),
        ("state_tax", "State/UT Tax Amount"),
        ("cess", "Cess Amount"),
    ],
    "hsn (b2c)": [
        ("hsn", "HSN"),
        ("description", "Description"),
        ("uqc", "UQC"),
        ("total_quantity", "Total Quantity"),
        ("total_value", "Total Value"),
        ("taxable_value", "Taxable Value"),
        ("rate", "Rate"),
        ("integrated_tax", "Integrated Tax Amount"),
        ("central_tax", "Central Tax Amount"),
        ("state_tax", "State/UT Tax Amount"),
        ("cess", "Cess Amount"),
    ],
    "b2cs": [
        ("type", "Type"),
        ("place_of_supply", "Place Of Supply"),
        ("applicable_tax_rate", "Applicable % of Tax Rate"),
        ("rate", "Rate"),
        ("taxable_value", "Taxable Value"),
        ("cess", "Cess Amount"),
        ("ecommerce_gstin", "E-Commerce GSTIN"),
    ],
    "docs": [
        ("nature_of_document", "Nature of Document"),
        ("sr_no_from", "Sr. No. From"),
        ("sr_no_to", "Sr. No. To"),
        ("total_number", "Total Number"),
        ("cancelled", "Cancelled"),
    ],
}

# excel_rows dict key (see gstr1_generator.build_gstr1_excel_rows) for each sheet
_SHEET_DATA_KEY = {
    "b2b": "b2b", "cdnr": "cdnr", "hsn (b2b)": "hsn_b2b",
    "hsn (b2c)": "hsn_b2c", "b2cs": "b2cs", "docs": "docs",
}

# Row keys whose value is a "DD-MM-YYYY" string (from gstr1_generator's
# _parse_date) that must be written as a real Excel date, matching the
# template's own pre-set numeric date format for that column -- a plain
# text string in a cell the template formats as a date is exactly what
# was confirmed causing ExpressGST to mis-read/reject the workbook.
_DATE_KEYS = {"invoice_date", "note_date"}


def _to_excel_date(value):
    if not value:
        return None
    try:
        return datetime.strptime(str(value), "%d-%m-%Y").date()
    except ValueError:
        return value  # already unparseable upstream; write as-is rather than drop it


def _header_column_map(ws) -> dict:
    """Maps each non-empty header label in the template's header row to its column index."""
    return {
        ws.cell(row=HEADER_ROW, column=c).value: c
        for c in range(1, ws.max_column + 1)
        if ws.cell(row=HEADER_ROW, column=c).value is not None
    }


def _write_sheet(wb, sheet_name: str, rows: List[dict]) -> None:
    columns = _SHEET_COLUMNS[sheet_name]
    ws = wb[sheet_name]
    header_map = _header_column_map(ws)

    resolved = []
    for key, label in columns:
        col_idx = header_map.get(label)
        if col_idx is None:
            raise ValueError(
                f"Column '{label}' not found in template sheet '{sheet_name}' — "
                f"the GSTN/ExpressGST template has likely changed; update _SHEET_COLUMNS."
            )
        resolved.append((key, col_idx))

    for row_offset, row in enumerate(rows):
        row_idx = DATA_START_ROW + row_offset
        for key, col_idx in resolved:
            value = row.get(key)
            if key in _DATE_KEYS:
                value = _to_excel_date(value)
            ws.cell(row=row_idx, column=col_idx, value=value)


def write_gstr1_excel(excel_rows: dict, path: str) -> None:
    """
    Writes the GSTN-offline-tool-format GSTR-1 workbook to `path`, based on
    the real reference template (see module docstring).

    Args:
        excel_rows: services.gstr1_generator.build_gstr1_excel_rows(items, firm_gstin)'s output.
        path: output .xlsx path.
    """
    if not OPENPYXL_AVAILABLE:
        raise RuntimeError("openpyxl not installed")

    wb = openpyxl.load_workbook(TEMPLATE_PATH)

    for sheet_name, data_key in _SHEET_DATA_KEY.items():
        _write_sheet(wb, sheet_name, excel_rows.get(data_key, []))

    wb.save(path)
    total_rows = sum(len(excel_rows.get(k, [])) for k in _SHEET_DATA_KEY.values())
    logger.info(f"Wrote GSTR-1 offline-tool workbook to {path} ({total_rows} rows across {len(_SHEET_DATA_KEY)} sheets)")
